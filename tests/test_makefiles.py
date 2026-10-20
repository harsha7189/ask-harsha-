import io
import json

import pytest

from agent import makefiles as mf

MD = "# Plan\n\nShip **export** by *Dec 1*.\n\n- Beta\n- GA\n\n1. One\n2. Two\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\n```\ncode()\n```\n"


def test_every_format_builds():
    pdf = mf.build({"filename": "plan.pdf", "title": "Plan", "content": MD})
    assert pdf.data.startswith(b"%PDF") and pdf.mime == "application/pdf"
    docx = mf.build({"filename": "plan.docx", "content": MD})
    assert docx.data[:2] == b"PK"
    from docx import Document
    text = "\n".join(p.text for p in Document(io.BytesIO(docx.data)).paragraphs)
    assert "Ship export by Dec 1." in text and "Beta" in text
    xlsx = mf.build({"filename": "b.xlsx", "sheets": [{"name": "Budget", "rows": [["Item", "Cost"], ["Desk", 300], ["Total", "=SUM(B2:B2)"]]}]})
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.data)).active
    assert ws.title == "Budget" and ws["B2"].value == 300 and ws["B3"].value == "=SUM(B2:B2)" and ws.freeze_panes == "A2"
    csv = mf.build({"filename": "b.csv", "sheets": [{"rows": [["a", "b"], [1, None]]}]})
    assert csv.data.decode("utf-8-sig") == "a,b\r\n1,\r\n"
    html = mf.build({"filename": "p.html", "content": "<h1>Hi</h1>"})
    assert b"<!doctype html>" in html.data and b"<h1>Hi</h1>" in html.data
    py = mf.build({"filename": "../../etc/app.py", "content": "print(1)\n"})
    assert py.filename == "app.py" and py.data == b"print(1)\n"            # no paths, just a name


@pytest.mark.parametrize("spec,msg", [
    ({"filename": "x.exe", "content": "a"}, "can't make .exe"),
    ({"filename": "noext", "content": "a"}, "extension"),
    ({"filename": "a.xlsx"}, "needs `sheets`"),
    ({"filename": "a.pdf", "content": "  "}, "needs `content`"),
    ({"filename": "a.txt", "content": "x" * (mf.MAX_CONTENT + 1)}, "too long"),
])
def test_bad_specs_explain_the_fix(spec, msg):
    with pytest.raises(mf.FileError, match=msg):
        mf.build(spec)


def test_the_agent_makes_a_file_and_the_page_can_download_it():
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient
    from agent import inbox as inboxmod, receipts as rcpt
    from agent.notes import MemoryNotesStore
    from fakes import RoutingClient, response, text, tool_use
    from server.app import Settings, create_app
    spec = {"filename": "budget.xlsx", "sheets": [{"name": "Budget", "rows": [["Item", "Cost"], ["Desk", 300]]}]}
    client = RoutingClient([response([tool_use("t1", "create_file", spec)], "tool_use"),
                            response([text("Here's your budget spreadsheet with one item.")], "end_turn")],
                           plan={"expects_file": True, "file_types": ["xlsx"], "checks": [], "cannot_check": []})
    s = Settings(daily_budget_usd=5.0, rate_limit=50, allowed_origins=[], owner_passcode="open-sesame")
    app = create_app(client=client, settings=s, notes=MemoryNotesStore(), inbox=inboxmod.MemoryInboxStore(),
                     receipts=rcpt.ReceiptStore(inboxmod.make_key("k")), accounts=False)
    c = TestClient(app)
    d = c.post("/api/chat", json={"question": "Make me a budget spreadsheet"}).json()
    assert d["files"][0]["filename"] == "budget.xlsx" and d["files"][0]["sheets"][0]["rows"] == 2
    assert d["proof"]["status"] == "verified" and d["proof"]["attempts"] == 1          # files are always verified
    assert client.requests[0]["max_tokens"] == 8000 and any(t["name"] == "create_file" for t in client.requests[0]["tools"])
    tool_reply = json.loads(client.requests[1]["messages"][-1]["content"][0]["content"])
    assert tool_reply["ok"] is True
    r = c.post("/api/files/build", json={"spec": d["files"][0]["spec"]})
    assert r.status_code == 200 and r.content[:2] == b"PK" and 'filename="budget.xlsx"' in r.headers["content-disposition"]
    assert c.post("/api/files/build", json={"spec": {"filename": "x.exe", "content": "a"}}).status_code == 400
