"""Files the agent makes for people: Excel, CSV, PDF, Word, HTML and code.

The model calls the `create_file` tool with a small spec: a filename, and either `content` (text, code, HTML, or
Markdown for PDF and Word) or `sheets` (rows for Excel and CSV). This module turns the spec into real bytes.

Nothing is stored on the server. The spec travels with the answer (and with saved chats), and the page asks
/api/files/build to turn it into the file whenever someone taps Download, so files never expire.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

FONTS = Path(__file__).parent / "fonts"
MAX_CONTENT = 300_000          # characters of text, code or Markdown
MAX_SHEETS = 10
MAX_ROWS = 5_000
MAX_COLS = 60
MAX_CELL = 5_000

TEXT_TYPES = {
    "txt": "text/plain", "md": "text/markdown", "py": "text/x-python", "js": "text/javascript", "ts": "text/plain",
    "jsx": "text/plain", "tsx": "text/plain", "json": "application/json", "css": "text/css", "sql": "text/plain",
    "sh": "text/x-shellscript", "yaml": "text/yaml", "yml": "text/yaml", "xml": "application/xml", "java": "text/plain",
    "c": "text/plain", "cpp": "text/plain", "h": "text/plain", "cs": "text/plain", "go": "text/plain", "rb": "text/plain",
    "php": "text/plain", "swift": "text/plain", "kt": "text/plain", "r": "text/plain", "ics": "text/calendar",
    "svg": "image/svg+xml", "toml": "text/plain", "ini": "text/plain",
}
BINARY_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
TABLE_TYPES = {"csv": "text/csv", "tsv": "text/tab-separated-values"}
HTML_TYPES = {"html": "text/html", "htm": "text/html"}
KINDS = {**TEXT_TYPES, **BINARY_TYPES, **TABLE_TYPES, **HTML_TYPES}

TOOL_SPEC = {
    "name": "create_file",
    "description": (
        "Make a downloadable file for the person: an Excel spreadsheet (.xlsx), CSV, PDF, Word document (.docx), HTML "
        "page, or code/text file (.py, .js, .json, .md, .sql, and more). Use it when they ask for a file, a spreadsheet, "
        "a document to download, or code longer than about 20 lines. For .xlsx/.csv give `sheets`; for everything else "
        "give `content`. PDF and Word take Markdown: # headings, paragraphs, - bullets, 1. numbered lists, **bold**, "
        "| pipe | tables |, and ``` code blocks. Excel cells starting with = are formulas (=SUM(B2:B10)). After it "
        "succeeds, tell them in a sentence or two what's in the file; don't paste the contents into the chat."),
    "input_schema": {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "Short name with extension, e.g. budget-2026.xlsx, notes.pdf, app.py"},
            "title": {"type": "string", "description": "Optional document title for PDF and Word"},
            "content": {"type": "string", "description": "The file's text: code, HTML, plain text, CSV, or Markdown for PDF/Word"},
            "sheets": {
                "type": "array",
                "description": "For .xlsx (one entry per sheet) or .csv (first sheet). The first row is the header.",
                "items": {"type": "object", "properties": {
                    "name": {"type": "string"},
                    "rows": {"type": "array", "items": {"type": "array", "items": {}}},
                }, "required": ["rows"]},
            },
        },
        "required": ["filename"],
    },
}


class FileError(ValueError):
    """The spec can't be made into a file; the message tells the model what to fix."""


@dataclass
class Built:
    filename: str
    mime: str
    data: bytes

    @property
    def kind(self) -> str:
        return self.filename.rsplit(".", 1)[-1].lower()


def clean_name(name: str) -> str:
    name = Path(str(name or "").strip().replace("\\", "/")).name
    name = re.sub(r"[^\w.\- ()]+", "-", name).strip(" .-")[:80]
    if "." not in name:
        raise FileError("The filename needs an extension, like report.pdf or budget.xlsx.")
    ext = name.rsplit(".", 1)[-1].lower()
    if ext not in KINDS:
        raise FileError(f"I can't make .{ext} files. Try one of: xlsx, csv, pdf, docx, html, md, txt, py, js, json, sql.")
    return name


def check(spec: dict) -> dict:
    """Validate and normalize a spec. Returns the clean spec (what's stored with the answer)."""
    if not isinstance(spec, dict):
        raise FileError("Send the file as an object with a filename.")
    name = clean_name(spec.get("filename", ""))
    ext = name.rsplit(".", 1)[-1].lower()
    out = {"filename": name}
    if spec.get("title"):
        out["title"] = " ".join(str(spec["title"]).split())[:200]
    sheets = spec.get("sheets")
    content = spec.get("content")
    if ext == "xlsx" or (ext in TABLE_TYPES and sheets):
        if not isinstance(sheets, list) or not sheets:
            raise FileError("An Excel file needs `sheets`: [{\"name\": \"Sheet1\", \"rows\": [[\"Header\", ...], [...]]}].")
        clean = []
        for i, sh in enumerate(sheets[:MAX_SHEETS]):
            rows = sh.get("rows") if isinstance(sh, dict) else None
            if not isinstance(rows, list) or not rows:
                raise FileError(f"Sheet {i + 1} has no rows.")
            if len(rows) > MAX_ROWS:
                raise FileError(f"Sheet {i + 1} has {len(rows)} rows; the limit is {MAX_ROWS}.")
            fixed = []
            for r in rows:
                r = r if isinstance(r, list) else [r]
                fixed.append([_cell(v) for v in r[:MAX_COLS]])
            title = re.sub(r"[\[\]:*?/\\]", "-", str((sh.get("name") if isinstance(sh, dict) else "") or f"Sheet{i + 1}"))[:31]
            clean.append({"name": title or f"Sheet{i + 1}", "rows": fixed})
        out["sheets"] = clean
    else:
        if not isinstance(content, str) or not content.strip():
            raise FileError("This file type needs `content` with the file's text.")
        if len(content) > MAX_CONTENT:
            raise FileError(f"That's too long ({len(content):,} characters). Keep it under {MAX_CONTENT:,}.")
        out["content"] = content
    return out


def _cell(v):
    if v is None or isinstance(v, (bool, int, float)):
        return v
    return str(v)[:MAX_CELL]


def build(spec: dict) -> Built:
    spec = check(spec)
    name = spec["filename"]
    ext = name.rsplit(".", 1)[-1].lower()
    if ext == "xlsx":
        return Built(name, BINARY_TYPES[ext], _xlsx(spec["sheets"]))
    if ext in TABLE_TYPES:
        if "sheets" in spec:
            buf = io.StringIO()
            w = csv.writer(buf, delimiter="\t" if ext == "tsv" else ",")
            for r in spec["sheets"][0]["rows"]:
                w.writerow(["" if v is None else v for v in r])
            text = buf.getvalue()
        else:
            text = spec["content"]
        return Built(name, TABLE_TYPES[ext] + "; charset=utf-8", ("﻿" + text).encode("utf-8"))  # BOM: Excel reads UTF-8
    if ext == "pdf":
        return Built(name, BINARY_TYPES[ext], _pdf(spec["content"], spec.get("title")))
    if ext == "docx":
        return Built(name, BINARY_TYPES[ext], _docx(spec["content"], spec.get("title")))
    if ext in HTML_TYPES:
        html = spec["content"]
        if "<html" not in html.lower():
            title = spec.get("title") or name.rsplit(".", 1)[0]
            html = (f"<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n<meta name=\"viewport\" "
                    f"content=\"width=device-width, initial-scale=1\">\n<title>{_esc(title)}</title>\n</head>\n<body>\n{html}\n</body>\n</html>\n")
        return Built(name, "text/html; charset=utf-8", html.encode("utf-8"))
    return Built(name, TEXT_TYPES[ext] + "; charset=utf-8", spec["content"].encode("utf-8"))


def summary(spec: dict, built: Built) -> dict:
    """What the page needs to show a file card (the spec goes along so Download can rebuild it)."""
    info = {"filename": built.filename, "kind": built.kind, "size": len(built.data)}
    if "sheets" in spec:
        info["sheets"] = [{"name": s["name"], "rows": len(s["rows"]), "cols": max((len(r) for r in s["rows"]), default=0)} for s in spec["sheets"]]
    else:
        info["lines"] = spec["content"].count("\n") + 1
    return info


# ---------- Excel ----------
def _xlsx(sheets: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")
    for sh in sheets:
        ws = wb.create_sheet(sh["name"])
        widths: dict[int, int] = {}
        for r_i, row in enumerate(sh["rows"], start=1):
            for c_i, v in enumerate(row, start=1):
                cell = ws.cell(row=r_i, column=c_i, value=v)
                if r_i == 1:
                    cell.font, cell.fill = head_font, head_fill
                    cell.alignment = Alignment(vertical="center")
                shown = len(str(v)) if v is not None and not (isinstance(v, str) and v.startswith("=")) else 10
                widths[c_i] = max(widths.get(c_i, 0), min(shown, 60))
        for c_i, w in widths.items():
            ws.column_dimensions[get_column_letter(c_i)].width = max(8, w + 2)
        if len(sh["rows"]) > 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- Markdown → PDF and Word ----------
def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _blocks(md: str) -> list[tuple]:
    """A small Markdown reader: headings, paragraphs, bullets, numbered items, tables, code blocks, rules."""
    lines = md.replace("\r\n", "\n").split("\n")
    out, para, i = [], [], 0

    def flush():
        if para:
            out.append(("p", " ".join(s.strip() for s in para)))
            para.clear()
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            code, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            out.append(("code", "\n".join(code)))
        elif m := re.match(r"^(#{1,4})\s+(.*)", s):
            flush()
            out.append(("h", len(m.group(1)), m.group(2).strip()))
        elif s.startswith("|") and s.endswith("|") and s.count("|") >= 2:
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            out.append(("table", rows))
            continue
        elif m := re.match(r"^[-*•]\s+(.*)", s):
            flush()
            out.append(("li", m.group(1)))
        elif m := re.match(r"^(\d+)[.)]\s+(.*)", s):
            flush()
            out.append(("ol", int(m.group(1)), m.group(2)))
        elif re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            flush()
            out.append(("hr",))
        elif not s:
            flush()
        else:
            para.append(s)
        i += 1
    flush()
    return out


def _inline_rl(t: str) -> str:
    """Markdown bold/italic/code → reportlab's mini markup."""
    t = _esc(t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"<i>\1</i>", t)
    t = re.sub(r"`([^`]+)`", r"<font face='DejaVuMono'>\1</font>", t)
    t = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r"<link href='\2' color='#1d4ed8'>\1</link>", t)
    return t


_fonts_ready = False


def _register_fonts():
    global _fonts_ready
    if _fonts_ready:
        return
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("DejaVu", str(FONTS / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONTS / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVuMono", str(FONTS / "DejaVuSansMono.ttf")))
    for italic in (0, 1):   # no italic face bundled: italic text uses the regular face
        addMapping("DejaVu", 0, italic, "DejaVu")
        addMapping("DejaVu", 1, italic, "DejaVu-Bold")
    _fonts_ready = True


def _pdf(md: str, title: str | None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (HRFlowable, ListFlowable, ListItem, Paragraph, Preformatted, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)
    _register_fonts()
    base = ParagraphStyle("base", fontName="DejaVu", fontSize=10.5, leading=15, alignment=TA_LEFT, spaceAfter=6)
    heads = {1: ParagraphStyle("h1", parent=base, fontName="DejaVu-Bold", fontSize=20, leading=25, spaceBefore=4, spaceAfter=10),
             2: ParagraphStyle("h2", parent=base, fontName="DejaVu-Bold", fontSize=15, leading=20, spaceBefore=10, spaceAfter=6),
             3: ParagraphStyle("h3", parent=base, fontName="DejaVu-Bold", fontSize=12.5, leading=17, spaceBefore=8, spaceAfter=4),
             4: ParagraphStyle("h4", parent=base, fontName="DejaVu-Bold", fontSize=11, leading=15, spaceBefore=6, spaceAfter=3)}
    cell = ParagraphStyle("cell", parent=base, fontSize=9.5, leading=12.5, spaceAfter=0)
    cell_b = ParagraphStyle("cellb", parent=cell, fontName="DejaVu-Bold")
    code = ParagraphStyle("code", parent=base, fontName="DejaVuMono", fontSize=9, leading=12, backColor=colors.HexColor("#f3f4f6"),
                          borderPadding=6, spaceBefore=4, spaceAfter=10)
    story, pending = [], []

    def close_list():
        if pending:
            kind = pending[0][0]
            items = [ListItem(Paragraph(_inline_rl(t), base), leftIndent=14) for _, t in pending]
            if kind == "ol":
                story.append(ListFlowable(items, bulletType="1", bulletFormat="%s.", leftIndent=18, bulletFontName="DejaVu", bulletFontSize=10))
            else:
                story.append(ListFlowable(items, bulletType="bullet", start="\u2022", leftIndent=16, bulletFontName="DejaVu", bulletFontSize=10))
            pending.clear()
    if title:
        story.append(Paragraph(_inline_rl(title), heads[1]))
    width = LETTER[0] - 1.6 * inch
    for b in _blocks(md):
        if b[0] in ("li", "ol"):
            if pending and pending[0][0] != b[0]:
                close_list()
            pending.append((b[0], b[-1]))
            continue
        close_list()
        if b[0] == "h":
            if title and b[1] == 1 and b[2].strip() == title.strip():
                continue
            story.append(Paragraph(_inline_rl(b[2]), heads[min(b[1], 4)]))
        elif b[0] == "p":
            story.append(Paragraph(_inline_rl(b[1]), base))
        elif b[0] == "code":
            story.append(Preformatted(b[1], code))
        elif b[0] == "hr":
            story.append(HRFlowable(width="100%", color=colors.HexColor("#d1d5db"), spaceBefore=6, spaceAfter=10))
        elif b[0] == "table" and b[1]:
            n = max(len(r) for r in b[1])
            rows = [[Paragraph(_inline_rl(c), cell_b if r_i == 0 else cell) for c in r + [""] * (n - len(r))] for r_i, r in enumerate(b[1])]
            t = Table(rows, colWidths=[width / n] * n, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
            story += [t, Spacer(1, 10)]
    close_list()
    if not story:
        story.append(Paragraph(" ", base))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("DejaVu", 8.5)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawRightString(LETTER[0] - 0.8 * inch, 0.5 * inch, f"{doc.page}")
        canvas.drawString(0.8 * inch, 0.5 * inch, "Made with Ask Harsha")
        canvas.restoreState()
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.8 * inch,
                      bottomMargin=0.8 * inch, title=title or "", author="Ask Harsha").build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _docx(md: str, title: str | None) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    def runs(par, text):
        for part in re.split(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*(?!\s)[^*]+?(?<!\s)\*(?![\w*]))", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                par.add_run(part[1:-1]).italic = True
            elif part.startswith("`") and part.endswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = "Consolas"
            else:
                par.add_run(re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r"\1 (\2)", part))
    if title:
        doc.add_heading(title, level=0)
    for b in _blocks(md):
        if b[0] == "h":
            if title and b[1] == 1 and b[2].strip() == title.strip():
                continue
            doc.add_heading(re.sub(r"\*\*(.+?)\*\*", r"\1", b[2]), level=min(b[1], 4))
        elif b[0] == "p":
            runs(doc.add_paragraph(), b[1])
        elif b[0] == "li":
            runs(doc.add_paragraph(style="List Bullet"), b[1])
        elif b[0] == "ol":
            runs(doc.add_paragraph(style="List Number"), b[2])
        elif b[0] == "code":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.font.name, r.font.size = "Consolas", Pt(9.5)
        elif b[0] == "hr":
            doc.add_paragraph("―" * 30)
        elif b[0] == "table" and b[1]:
            n = max(len(r) for r in b[1])
            t = doc.add_table(rows=0, cols=n)
            t.style = "Table Grid"
            for r_i, r in enumerate(b[1]):
                cells = t.add_row().cells
                for c_i in range(n):
                    par = cells[c_i].paragraphs[0]
                    text = r[c_i] if c_i < len(r) else ""
                    if r_i == 0:
                        par.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", text)).bold = True
                    else:
                        runs(par, text)
            doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
