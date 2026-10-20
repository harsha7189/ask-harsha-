"""Proof Mode: never confuse trying with succeeding.

For a task (make a file, write a report), Proof Mode:
  1. captures intent: turns the person's request into a checklist code can verify. A small model writes the checklist
     from the request alone, before it sees any answer, so the work can't grade itself.
  2. lets the agent do the work.
  3. verifies the real result with deterministic code: rebuilds each file, opens it with an independent reader
     (openpyxl, pypdf, python-docx, Python's compiler, json, html.parser), and checks it against the checklist.
     No model is asked whether it succeeded.
  4. if something failed and can be fixed, tells the agent exactly what failed and lets it try again (at most twice).
  5. returns a completion report: every check with its evidence, the states with timestamps, and a final status:
     Verified complete, Partially complete, Failed or Unverified. The server signs it with the Ed25519 key.

The signature proves the report wasn't altered after Ask Harsha issued it. It is not proof of anything that happened
outside Ask Harsha.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import secrets
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from html.parser import HTMLParser
from typing import Any

from . import makefiles, speed

MAX_ATTEMPTS = 3          # the first try plus two fixes
CHECK_TYPES = ("columns", "min_rows", "sections", "mentions", "max_pages", "min_words", "max_words", "min_items")
FILE_TYPES = ("xlsx", "csv", "pdf", "docx", "html", "md", "txt", "py", "js", "json", "sql")

TASK_WORDS = re.compile(
    r"\b(make|create|build|generate|write|draft|prepare|produce|put together|design|export|convert|turn .* into|give me (a|an)|"
    r"spreadsheet|excel|xlsx|csv|pdf|word doc|docx|document|report|template|tracker|checklist|plan|itinerary|resume|"
    r"cover letter|landing page|web ?page|html|script|code|python|json)\b", re.I)

PLAN_PROMPT = """You turn a person's request into a checklist that code will use to verify the result. You only see the \
request, never the answer. Call set_checklist once.

- expects_file: true if they asked for a file to download (spreadsheet, Excel, CSV, PDF, Word, HTML page, code file).
- file_types: the formats they asked for: xlsx for Excel or "spreadsheet", csv, pdf, docx for Word, html for a web page, \
py/js/json/sql/md/txt for code or text files. Empty if no file.
- checks: only requirements the request clearly states or plainly implies. Types:
  columns (list of column names a table must have), min_rows (at least N data rows), sections (headings that must \
appear), mentions (exact words or names that must appear), max_pages (a PDF page limit, e.g. "one page" = 1), min_words, \
max_words, min_items (at least N list items, e.g. "5 tips"). Give each a short plain label, like "Has 5 jobs".
- cannot_check: expectations code can't verify, like tone, accuracy of facts or quality. Keep each under 10 words.
Don't invent requirements the person didn't ask for."""

PLAN_TOOL = {
    "name": "set_checklist",
    "description": "The checklist code will verify.",
    "input_schema": {
        "type": "object",
        "properties": {
            "expects_file": {"type": "boolean"},
            "file_types": {"type": "array", "items": {"type": "string", "enum": list(FILE_TYPES)}},
            "checks": {"type": "array", "items": {"type": "object", "properties": {
                "type": {"type": "string", "enum": list(CHECK_TYPES)},
                "value": {"description": "List of strings for columns/sections/mentions; a number otherwise."},
                "label": {"type": "string"}}, "required": ["type", "value", "label"]}},
            "cannot_check": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["expects_file", "file_types", "checks", "cannot_check"],
    },
}

KNOWN_FUNCTIONS = set("""SUM AVERAGE MIN MAX COUNT COUNTA COUNTBLANK IF IFS IFERROR AND OR NOT ROUND ROUNDUP ROUNDDOWN INT MOD ABS
SQRT POWER PRODUCT SUMPRODUCT SUMIF SUMIFS COUNTIF COUNTIFS AVERAGEIF AVERAGEIFS MAXIFS MINIFS VLOOKUP HLOOKUP XLOOKUP INDEX
MATCH XMATCH CHOOSE TODAY NOW DATE YEAR MONTH DAY WEEKDAY EDATE EOMONTH DATEDIF NETWORKDAYS WORKDAY TEXT VALUE CONCAT
CONCATENATE TEXTJOIN LEFT RIGHT MID LEN UPPER LOWER PROPER TRIM SUBSTITUTE REPLACE FIND SEARCH PMT FV PV NPV IRR RATE NPER
MEDIAN MODE STDEV STDEV.S STDEV.P VAR LARGE SMALL RANK RANK.EQ PERCENTILE QUARTILE FILTER SORT SORTBY UNIQUE SEQUENCE LET
TRANSPOSE HYPERLINK ISBLANK ISNUMBER ISTEXT ISERROR OFFSET INDIRECT ROW ROWS COLUMN COLUMNS CEILING FLOOR SIGN EXP LN LOG
LOG10 PI RAND RANDBETWEEN TRUE FALSE N T""".split())
PLACEHOLDER = re.compile(r"lorem ipsum|\[(insert|your|add|company name|name)[^\]]*\]|\{\{\s*\w+\s*\}\}|<placeholder>|\bTBD\b|\bXXX+\b", re.I)
NUMWORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
            "twelve": 12, "fifteen": 15, "twenty": 20}
CLAIMS_DONE = re.compile(r"\b(here('s| is| are)|i('ve| have) (created|made|built|prepared|put together|written|added)|done|ready|"
                         r"attached|all set|completed?)\b", re.I)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def looks_like_task(question: str) -> bool:
    return bool(TASK_WORDS.search(question or ""))


# ---------------------------------------------------------------- the plan (what success means)
@dataclass
class Plan:
    expects_file: bool = False
    file_types: list[str] = field(default_factory=list)
    checks: list[dict] = field(default_factory=list)
    cannot_check: list[str] = field(default_factory=list)
    source: str = "rules"

    def to_dict(self) -> dict:
        return {"expects_file": self.expects_file, "file_types": self.file_types, "checks": self.checks,
                "cannot_check": self.cannot_check, "source": self.source}


def _num(v) -> int | None:
    try:
        n = int(float(v))
        return n if 0 < n <= 100_000 else None
    except (TypeError, ValueError):
        return None


def clean_plan(raw: dict, source: str) -> Plan:
    p = Plan(source=source)
    p.expects_file = bool(raw.get("expects_file"))
    p.file_types = [t for t in (raw.get("file_types") or []) if t in FILE_TYPES][:5]
    for c in (raw.get("checks") or [])[:12]:
        if not isinstance(c, dict) or c.get("type") not in CHECK_TYPES:
            continue
        t, v = c["type"], c.get("value")
        if t in ("columns", "sections", "mentions"):
            vals = v if isinstance(v, list) else [v]
            vals = [" ".join(str(x).split())[:80] for x in vals if str(x).strip()][:15]
            if not vals:
                continue
            v = vals
        else:
            v = _num(v)
            if v is None:
                continue
        label = " ".join(str(c.get("label") or "").split())[:90] or _default_label(t, v)
        p.checks.append({"type": t, "value": v, "label": label})
    p.cannot_check = [" ".join(str(x).split())[:90] for x in (raw.get("cannot_check") or []) if str(x).strip()][:5]
    if p.file_types:
        p.expects_file = True
    return p


def _default_label(t: str, v) -> str:
    return {"columns": f"Has columns: {', '.join(v) if isinstance(v, list) else v}",
            "min_rows": f"At least {v} rows", "sections": f"Has sections: {', '.join(v) if isinstance(v, list) else v}",
            "mentions": f"Mentions {', '.join(v) if isinstance(v, list) else v}", "max_pages": f"At most {v} page{'s' if v != 1 else ''}",
            "min_words": f"At least {v} words", "max_words": f"At most {v} words", "min_items": f"At least {v} items"}[t]


def rules_plan(request: str) -> Plan:
    """A plan from simple rules, used when the model isn't available."""
    q = request.lower()
    types = []
    for t, pat in (("xlsx", r"\b(excel|xlsx|spreadsheet)\b"), ("csv", r"\bcsv\b"), ("pdf", r"\bpdf\b"), ("docx", r"\b(word doc|docx|word file)\b"),
                   ("html", r"\b(html|web ?page|landing page)\b"), ("py", r"\b(python|\.py)\b"), ("json", r"\bjson\b")):
        if re.search(pat, q):
            types.append(t)
    checks = []
    m = re.search(r"\b(\d{1,3}|" + "|".join(NUMWORDS) + r")\s+(\w+\s+)?(jobs|rows|items|tips|ideas|steps|questions|companies|"
                  r"products|tasks|examples|reasons|options|people|books|places|names|recipes)\b", q)
    if m:
        n = int(m.group(1)) if m.group(1).isdigit() else NUMWORDS[m.group(1)]
        checks.append({"type": "min_rows" if "xlsx" in types or "csv" in types else "min_items", "value": n,
                       "label": f"Has at least {n} {m.group(3)}"})
    if re.search(r"\bone[- ]page\b", q) and "pdf" in types:
        checks.append({"type": "max_pages", "value": 1, "label": "Fits on one page"})
    return Plan(expects_file=bool(types), file_types=types, checks=checks, source="rules")


def capture(client: Any, model: str, request: str) -> tuple[Plan, int, int]:
    """Intent capture: the checklist for this request. Falls back to rules if the model call fails."""
    try:
        resp = client.messages.create(model=model, max_tokens=600, system=PLAN_PROMPT, tools=[PLAN_TOOL],
                                      tool_choice={"type": "tool", "name": "set_checklist"}, **speed.params(model, "helper"),
                                      messages=[{"role": "user", "content": request[:4000]}])
    except Exception:
        return rules_plan(request), 0, 0
    usage = _get(resp, "usage")
    i = (_get(usage, "input_tokens", 0) or 0) if usage is not None else 0
    o = (_get(usage, "output_tokens", 0) or 0) if usage is not None else 0
    block = next((b for b in (_get(resp, "content", []) or []) if _get(b, "type") == "tool_use"), None)
    raw = dict(_get(block, "input", {}) or {}) if block is not None else {}
    if not raw:
        return rules_plan(request), i, o
    plan = clean_plan(raw, "model")
    rp = rules_plan(request)   # formats named in plain words are never dropped
    for t in rp.file_types:
        if t not in plan.file_types:
            plan.file_types.append(t)
            plan.expects_file = True
    return plan, i, o


# ---------------------------------------------------------------- inspecting the real result
@dataclass
class Inspection:
    name: str
    kind: str
    opened: bool = False
    error: str = ""
    reader: str = ""
    sha256: str = ""
    size: int = 0
    text: str = ""
    headings: list[str] = field(default_factory=list)
    items: int = 0
    pages: int | None = None
    tables: list[dict] = field(default_factory=list)     # [{name, header, rows (data rows)}]
    problems: list[str] = field(default_factory=list)    # formula and syntax problems found while reading

    @property
    def words(self) -> int:
        return len(re.findall(r"\b\w+\b", self.text))


class _HTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.text, self.headings, self.items, self._h, self._skip = [], [], 0, None, 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        if re.fullmatch(r"h[1-6]", tag):
            self._h = []
        if tag == "li":
            self.items += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1
        if re.fullmatch(r"h[1-6]", tag) and self._h is not None:
            self.headings.append(" ".join("".join(self._h).split()))
            self._h = None

    def handle_data(self, data):
        if self._skip:
            return
        self.text.append(data)
        if self._h is not None:
            self._h.append(data)


def _md_shape(text: str) -> tuple[list[str], int]:
    heads, items = [], 0
    for line in text.splitlines():
        s = line.strip()
        if m := re.match(r"^#{1,6}\s+(.*)", s):
            heads.append(m.group(1).strip("* "))
        elif re.fullmatch(r"\*\*[^*]{2,80}\*\*:?", s) or (s.endswith(":") and 2 < len(s) < 60 and len(s.split()) <= 6):
            heads.append(s.strip("*: "))
        if re.match(r"^([-*•]|\d+[.)])\s+\S", s):
            items += 1
    return heads, items


def inspect_file(f: dict) -> Inspection:
    """Rebuild the file from its spec and open the bytes with an independent reader."""
    spec = f.get("spec") or {}
    name = str(spec.get("filename") or f.get("filename") or "file")
    ins = Inspection(name=name, kind=name.rsplit(".", 1)[-1].lower())
    try:
        built = makefiles.build(spec)
    except Exception as e:
        ins.error = f"couldn't be built: {e}"
        return ins
    data = built.data
    ins.sha256, ins.size = hashlib.sha256(data).hexdigest(), len(data)
    try:
        if ins.kind == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data))
            ins.reader = "openpyxl"
            texts = []
            for ws in wb.worksheets:
                rows = [[c.value for c in r] for r in ws.iter_rows()]
                rows = [r for r in rows if any(v not in (None, "") for v in r)]
                header = [str(v) if v is not None else "" for v in (rows[0] if rows else [])]
                ins.tables.append({"name": ws.title, "header": header, "rows": max(0, len(rows) - 1)})
                texts += [" ".join(str(v) for v in r if v is not None) for r in rows]
                ins.problems += _formula_problems(ws, wb)
            ins.text = "\n".join(texts)
            ins.items = max((t["rows"] for t in ins.tables), default=0)
        elif ins.kind in ("csv", "tsv"):
            rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter="\t" if ins.kind == "tsv" else ","))
            rows = [r for r in rows if any(c.strip() for c in r)]
            ins.reader = "csv module"
            ins.tables.append({"name": name, "header": rows[0] if rows else [], "rows": max(0, len(rows) - 1)})
            ins.text = "\n".join(" ".join(r) for r in rows)
            ins.items = ins.tables[0]["rows"]
        elif ins.kind == "pdf":
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            ins.reader = "pypdf"
            ins.pages = len(reader.pages)
            ins.text = "\n".join((p.extract_text() or "") for p in reader.pages)
            ins.headings = [ln.strip() for ln in ins.text.splitlines() if 0 < len(ln.strip()) <= 70]
            ins.items = len(re.findall(r"(?m)^\s*(•|\d+\.)\s+\S", ins.text))
        elif ins.kind == "docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            ins.reader = "python-docx"
            paras = doc.paragraphs
            ins.headings = [p.text.strip() for p in paras if (p.style.name or "").startswith(("Heading", "Title")) and p.text.strip()]
            ins.items = sum(1 for p in paras if (p.style.name or "").startswith("List"))
            ins.text = "\n".join(p.text for p in paras) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            for t in doc.tables:
                cells = [[c.text.strip() for c in r.cells] for r in t.rows]
                ins.tables.append({"name": "table", "header": cells[0] if cells else [], "rows": max(0, len(cells) - 1)})
        elif ins.kind in ("html", "htm", "svg"):
            p = _HTML()
            p.feed(data.decode("utf-8", "replace"))
            ins.reader = "html.parser"
            ins.text, ins.headings, ins.items = " ".join(" ".join(p.text).split()), p.headings, p.items
        else:
            src = data.decode("utf-8")
            ins.text = src
            if ins.kind == "py":
                ins.reader = "Python compiler (not run)"
                try:
                    compile(src, name, "exec")
                except SyntaxError as e:
                    ins.problems.append(f"Python syntax error on line {e.lineno}: {e.msg}")
            elif ins.kind == "json":
                ins.reader = "json parser"
                try:
                    json.loads(src)
                except json.JSONDecodeError as e:
                    ins.problems.append(f"JSON error on line {e.lineno}: {e.msg}")
            else:
                ins.reader = "text reader"
                ins.problems += _bracket_problems(src) if ins.kind in ("js", "ts", "css", "java", "c", "cpp", "go", "cs") else []
            ins.headings, ins.items = _md_shape(src) if ins.kind in ("md", "txt") else ([], 0)
        ins.opened = True
    except Exception as e:
        ins.error = f"{type(e).__name__}: {str(e)[:160]}"
    return ins


def inspect_answer(text: str) -> Inspection:
    ins = Inspection(name="the reply", kind="answer", opened=True, reader="the reply's text", text=text or "")
    ins.headings, ins.items = _md_shape(ins.text)
    ins.sha256, ins.size = hashlib.sha256(ins.text.encode()).hexdigest(), len(ins.text.encode())
    return ins


_REF = re.compile(r"(?:(?:'([^']+)'|([A-Za-z_][\w.]*))!)?\$?([A-Z]{1,3})\$?(\d{1,7})(?::\$?([A-Z]{1,3})\$?(\d{1,7}))?")


def _col_index(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


def _formula_problems(ws, wb) -> list[str]:
    out = []
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            f = v[1:]
            bare = re.sub(r'"[^"]*"', '""', f)
            if bare.count("(") != bare.count(")"):
                out.append(f"{ws.title}!{c.coordinate}: unbalanced parentheses in {v}")
                continue
            for fn in re.findall(r"([A-Z][A-Z0-9.]*)\s*\(", bare):
                if fn not in KNOWN_FUNCTIONS:
                    out.append(f"{ws.title}!{c.coordinate}: unknown function {fn}")
            for m in _REF.finditer(bare):
                sheet = m.group(1) or m.group(2)
                target = wb[sheet] if sheet and sheet in wb.sheetnames else ws
                if sheet and sheet not in wb.sheetnames:
                    out.append(f"{ws.title}!{c.coordinate}: refers to a sheet that doesn't exist ({sheet})")
                    continue
                cols = [m.group(3)] + ([m.group(5)] if m.group(5) else [])
                rows = [int(m.group(4))] + ([int(m.group(6))] if m.group(6) else [])
                if max(rows) > max(target.max_row, 1) or max(_col_index(x) for x in cols) > max(target.max_column, 1):
                    out.append(f"{ws.title}!{c.coordinate}: refers to {m.group(0)}, outside the data")
                    continue
                if re.match(r"(SUM|AVERAGE|MIN|MAX|PRODUCT)\(", bare) and len(rows) == 2:   # numbers typed as text are skipped by SUM
                    for r in range(min(rows), max(rows) + 1):
                        for ci in range(_col_index(cols[0]), _col_index(cols[-1]) + 1):
                            cell = target.cell(row=r, column=ci).value
                            if isinstance(cell, str) and re.fullmatch(r"\s*[$€£₹]?\s*-?[\d,]+(\.\d+)?\s*%?\s*", cell):
                                out.append(f"{target.title}!{target.cell(row=r, column=ci).coordinate} is the text '{cell}', "
                                           f"so {c.coordinate}'s formula ignores it (store it as a number)")
    return out[:20]


def _bracket_problems(src: str) -> list[str]:
    s = re.sub(r"(\"(\\.|[^\"\\])*\"|'(\\.|[^'\\])*'|`[^`]*`|//[^\n]*|/\*.*?\*/)", "", src, flags=re.S)
    pairs = {")": "(", "]": "[", "}": "{"}
    stack = []
    for ch in s:
        if ch in "([{":
            stack.append(ch)
        elif ch in pairs:
            if not stack or stack[-1] != pairs[ch]:
                return [f"unbalanced '{ch}'"]
            stack.pop()
    return [f"{len(stack)} unclosed bracket(s)"] if stack else []


# ---------------------------------------------------------------- verification
def _norm(s: str) -> str:
    return " ".join(re.sub(r"[^\w\s]", " ", str(s).lower()).split())


def _has(haystack: list[str], needle: str) -> bool:
    n = _norm(needle)
    return any(n and (n in _norm(h) or (_norm(h) and _norm(h) in n and len(_norm(h)) > 3)) for h in haystack)


@dataclass
class Result:
    status: str
    checks: list[dict]
    inspections: list[Inspection]
    fixable: list[dict]

    @property
    def counts(self) -> dict:
        c = {"pass": 0, "fail": 0, "unverified": 0}
        for ch in self.checks:
            c[ch["result"]] += 1
        return c


def verify(plan: Plan, files: list[dict], answer: str) -> Result:
    checks: list[dict] = []

    def add(label, result, evidence, target="", fixable=False):
        checks.append({"id": f"c{len(checks) + 1}", "label": label, "result": result, "evidence": evidence, "target": target,
                       "fixable": fixable and result == "fail"})
    files = list(files)
    ins = [inspect_file(f) for f in files]
    answer_ins = inspect_answer(answer)

    # 1. the files exist, are the asked-for kinds, and open
    if plan.expects_file and not files:
        add("A file was made", "fail", "No file was created in this answer.", fixable=True)
    for t in plan.file_types:
        got = [i for i in ins if i.kind == t or (t == "html" and i.kind == "htm")]
        if files:
            add(f"Made a .{t} file", "pass" if got else "fail",
                f"Found {got[0].name}." if got else f"Made {', '.join(i.name for i in ins)}, but no .{t} file.", fixable=not got)
    for i in ins:
        if i.opened:
            shape = (f"{len(i.tables)} sheet(s), " + ", ".join(f"{t['name']}: {t['rows']} rows × {len(t['header'])} columns" for t in i.tables)) if i.kind == "xlsx" \
                else (f"{i.pages} page(s), {i.words} words" if i.kind == "pdf" else f"{i.words} words" if i.kind in ("docx", "html", "md", "txt") else f"{i.text.count(chr(10)) + 1} lines")
            add(f"{i.name} opens", "pass", f"Opened with {i.reader}: {shape}. SHA-256 {i.sha256[:12]}…, {i.size:,} bytes.", i.name)
        else:
            add(f"{i.name} opens", "fail", f"It {i.error or 'could not be opened'}.", i.name, fixable=True)
            continue
        if i.problems:
            add(f"{i.name} has no errors", "fail", "; ".join(i.problems[:4]) + (f" (+{len(i.problems) - 4} more)" if len(i.problems) > 4 else ""),
                i.name, fixable=True)
        elif i.kind in ("xlsx", "py", "json", "js", "ts", "css"):
            what = "formulas refer to real cells and known functions" if i.kind == "xlsx" else "the syntax is valid (the code wasn't run)"
            add(f"{i.name} has no errors", "pass", f"Checked: {what}.", i.name)
        ph = sorted({m.group(0) for m in PLACEHOLDER.finditer(i.text)})
        if ph:
            add(f"{i.name} has no placeholders", "fail", "Still contains: " + ", ".join(f"“{p}”" for p in ph[:4]), i.name, fixable=True)

    # 2. what the request asked for, checked on the matching file (or on the reply when there's no file)
    targets = [i for i in ins if i.opened] or [answer_ins]
    for c in plan.checks:
        t, v, label = c["type"], c["value"], c["label"]
        if t in ("columns",):
            pool = [x for x in targets if x.tables]
            if not pool:
                add(label, "fail", "There's no table to check.", fixable=True)
                continue
            best = max(pool, key=lambda x: max(sum(_has(tb["header"], col) for col in v) for tb in x.tables))
            tb = max(best.tables, key=lambda tb: sum(_has(tb["header"], col) for col in v))
            missing = [col for col in v if not _has(tb["header"], col)]
            add(label, "fail" if missing else "pass",
                (f"Missing {', '.join(missing)}. " if missing else "") + f"Header in {best.name} ({tb['name']}): {', '.join(h for h in tb['header'] if h) or '(empty)'}.",
                best.name, fixable=bool(missing))
        elif t in ("min_rows", "min_items"):
            best = max(targets, key=lambda x: max([tb["rows"] for tb in x.tables] + [x.items]))
            n = max([tb["rows"] for tb in best.tables] + [best.items])
            unit = "data rows" if best.tables else "list items"
            add(label, "pass" if n >= v else "fail", f"Counted {n} {unit} in {best.name} (needed {v}).", best.name, fixable=n < v)
        elif t == "sections":
            best = max(targets, key=lambda x: sum(_has(x.headings + [x.text], s) for s in v))
            missing = [s for s in v if not _has(best.headings, s) and _norm(s) not in _norm(best.text)]
            add(label, "fail" if missing else "pass",
                f"Missing: {', '.join(missing)}." if missing else f"Found in {best.name}: {', '.join(v)}.", best.name, fixable=bool(missing))
        elif t == "mentions":
            best = max(targets, key=lambda x: sum(_norm(m) in _norm(x.text) for m in v))
            missing = [m for m in v if _norm(m) not in _norm(best.text)]
            add(label, "fail" if missing else "pass",
                f"Doesn't mention: {', '.join(missing)}." if missing else f"{best.name} mentions {', '.join(v)}.", best.name, fixable=bool(missing))
        elif t == "max_pages":
            pdfs = [x for x in targets if x.pages is not None]
            if not pdfs:
                add(label, "unverified", "No PDF to count pages in.")
                continue
            x = pdfs[0]
            add(label, "pass" if x.pages <= v else "fail", f"{x.name} has {x.pages} page(s) (limit {v}).", x.name, fixable=x.pages > v)
        elif t in ("min_words", "max_words"):
            x = max(targets, key=lambda x: x.words)
            ok = x.words >= v if t == "min_words" else x.words <= v
            add(label, "pass" if ok else "fail", f"{x.name} has {x.words:,} words ({'at least' if t == 'min_words' else 'at most'} {v:,} needed).",
                x.name, fixable=not ok)

    # 3. the reply doesn't claim files that weren't made
    names = {i.name.lower() for i in ins}
    claimed = {m.group(0).lower() for m in re.finditer(r"\b[\w\-]+\.(xlsx|csv|pdf|docx|html|py|js|json|md|txt|sql)\b", answer or "", re.I)}
    phantom = sorted(c for c in claimed if c not in names)
    if phantom:
        add("The reply only claims files that exist", "fail", f"The reply mentions {', '.join(phantom)}, but no such file was made.", fixable=True)
    elif claimed:
        add("The reply only claims files that exist", "pass", f"Every file the reply names exists: {', '.join(sorted(claimed))}.")

    # 4. what code can't judge, said plainly
    for item in plan.cannot_check:
        add(item, "unverified", "Code can't check this. Have a look yourself.")

    counts = {"pass": 0, "fail": 0, "unverified": 0}
    for ch in checks:
        counts[ch["result"]] += 1
    no_good_file = plan.expects_file and not any(i.opened for i in ins)
    if no_good_file or (counts["fail"] and not counts["pass"]):
        status = "failed"
    elif counts["fail"]:
        status = "partial"
    elif counts["pass"]:
        status = "verified"
    else:
        status = "unverified"
    return Result(status, checks, ins, [c for c in checks if c["fixable"]])


def fix_prompt(result: Result, attempt: int) -> str:
    lines = "\n".join(f"- {c['label']}: {c['evidence']}" for c in result.fixable)
    return ("[Proof Mode check, not from the person] Code checked your work against their request and found problems:\n"
            f"{lines}\n\nFix every one. If a file needs changes, call create_file again with the complete corrected file "
            "(same filename). Then write your reply to the person again from the start, as if for the first time, and say "
            f"plainly what's in the file. This is fix attempt {attempt - 1} of {MAX_ATTEMPTS - 1}.")


def claims_success(text: str) -> bool:
    return bool(CLAIMS_DONE.search(text or ""))


STATUS_WORDS = {"verified": "Verified complete", "partial": "Partially complete", "failed": "Failed", "unverified": "Unverified"}


def receipt(task_id: str, request: str, plan: Plan, result: Result, states: list[dict], attempts: int, model: str, key: str) -> dict:
    """What gets signed. It records what was checked and found; it can't prove events outside Ask Harsha."""
    return {
        "v": 1, "type": "proof", "id": task_id, "issued": now_iso(), "model": model, "key": key,
        "request": request[:600], "status": result.status, "attempts": attempts,
        "checks": [{"label": c["label"], "result": c["result"], "evidence": c["evidence"][:400]} for c in result.checks],
        "files": [{"name": i.name, "sha256": i.sha256, "bytes": i.size} for i in result.inspections if i.sha256],
        "states": states,
        "scope": "This signature proves Ask Harsha issued this record and it hasn't been altered. It is not proof of events outside Ask Harsha.",
    }


def new_task_id() -> str:
    return "P" + secrets.token_urlsafe(8)


def _get(obj: Any, key: str, default: Any = None) -> Any:
    if isinstance(obj, dict):
        return obj.get(key, default)
    return getattr(obj, key, default)


class Stats:
    """Proof Mode's own scoreboard: no content, just outcomes (owner can read it)."""

    def __init__(self):
        self.runs: list[dict] = []

    def add(self, run: dict) -> None:
        self.runs.append(run)
        del self.runs[:-2000]

    def summary(self, runs: list[dict] | None = None) -> dict:
        runs = self.runs if runs is None else runs
        n = len(runs)
        by = {s: sum(1 for r in runs if r["status"] == s) for s in STATUS_WORDS}
        retried = [r for r in runs if r["attempts"] > 1]
        fixed = [r for r in retried if r["status"] == "verified"]
        return {"tasks": n, "by_status": by,
                "verified_rate": round(by["verified"] / n, 3) if n else None,
                "false_success_caught": sum(1 for r in runs if r.get("caught")),
                "retried": len(retried), "fixed_by_retry": len(fixed),
                "retry_success_rate": round(len(fixed) / len(retried), 3) if retried else None,
                "avg_seconds": round(sum(r["seconds"] for r in runs) / n, 2) if n else None,
                "needed_person": sum(1 for r in runs if r["status"] != "verified")}
