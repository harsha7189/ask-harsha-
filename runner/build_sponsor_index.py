#!/usr/bin/env python3
"""Builds Job Pilot's H-1B sponsor index from the Department of Labor's public LCA disclosure data, and uploads it.

1. Download the "LCA Programs (H-1B, H-1B1, E-3)" disclosure files for the last two or three fiscal years from
   https://www.dol.gov/agencies/eta/foreign-labor/performance (Excel files, about 200 MB each).
2. python build_sponsor_index.py LCA_Disclosure_Data_FY2025_Q4.xlsx LCA_Disclosure_Data_FY2026_Q3.xlsx --upload

It keeps, per employer, only three numbers: certified H-1B filings for product roles, certified filings overall, and the
fiscal years they appear in. Nothing about any person is kept. Upload uses the runner token from `job_pilot.py setup`.

Needs: pip install openpyxl (only for .xlsx; .csv works without it)
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
try:
    from agent.jobsearch import norm_company   # same normalization the server uses for lookups
except Exception:   # running from a copy of runner/ without the rest of the repo
    def norm_company(name: str) -> str:
        n = re.sub(r"[^a-z0-9 ]", " ", (name or "").lower())
        n = re.sub(r"\b(inc|incorporated|llc|l l c|ltd|limited|corp|corporation|co|company|plc|lp|llp|holdings|group|technologies|technology|the)\b", " ", n)
        return " ".join(n.split())

PRODUCT_ROLE = re.compile(r"\bproduct\s+(manager|owner|management|lead)\b|\bproduct\s+mgr\b", re.I)
NOT_PRODUCT = re.compile(r"\bproduction\b|\bproduct\s+(design(er)?|marketing|engineer|support|specialist|analyst)\b", re.I)


def fiscal_year(path: str) -> str:
    m = re.search(r"FY\s?(\d{4})", path, re.I)
    return m.group(1) if m else ""


def rows(path: str):
    """Yields dicts with upper-case column names, streaming (the files are large)."""
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8", errors="replace") as fh:
            for r in csv.DictReader(fh):
                yield {k.strip().upper(): v for k, v in r.items() if k}
        return
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [str(h or "").strip().upper() for h in next(it)]
    for r in it:
        yield dict(zip(head, r))
    wb.close()


def build(paths: list[str]) -> dict:
    index: dict[str, dict] = {}
    for path in paths:
        fy = fiscal_year(path)
        n = 0
        for r in rows(path):
            status = str(r.get("CASE_STATUS") or "").lower()
            visa = str(r.get("VISA_CLASS") or "H-1B").upper()
            if not status.startswith("certified") or "H-1B" not in visa:     # "Certified" and "Certified - Withdrawn"
                continue
            key = norm_company(str(r.get("EMPLOYER_NAME") or ""))
            if not key:
                continue
            e = index.setdefault(key, {"pm": 0, "total": 0, "years": []})
            e["total"] += 1
            title = str(r.get("JOB_TITLE") or "")
            if PRODUCT_ROLE.search(title) and not NOT_PRODUCT.search(title):
                e["pm"] += 1
            if fy and fy not in e["years"]:
                e["years"].append(fy)
            n += 1
        print(f"{path}: {n:,} certified H-1B filings", file=sys.stderr)
    for e in index.values():
        e["years"] = sorted(e["years"])[-4:]
    return index


def upload(index: dict, source: str) -> dict:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from job_pilot import Keys, load_config
    cfg, token = load_config(), Keys().get("runner_token")
    if not cfg.get("server") or not token:
        sys.exit("Run `python job_pilot.py setup` first (it saves the server address and runner token).")
    req = urllib.request.Request(cfg["server"].rstrip("/") + "/api/jobs/sponsors", method="POST",
                                 data=json.dumps({"index": index, "source": source}).encode(),
                                 headers={"X-Runner-Token": token, "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=300) as r:
        return json.loads(r.read().decode())


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description="Build the H-1B sponsor index from DOL LCA disclosure files")
    ap.add_argument("files", nargs="+", help=".xlsx or .csv disclosure files")
    ap.add_argument("--out", default="sponsor_index.json")
    ap.add_argument("--upload", action="store_true", help="send it to Ask Harsha")
    args = ap.parse_args(argv)
    index = build(args.files)
    Path(args.out).write_text(json.dumps(index))
    product = sum(1 for e in index.values() if e["pm"])
    print(f"{len(index):,} employers, {product:,} with product-role filings → {args.out}")
    if args.upload:
        print(upload(index, ", ".join(Path(f).name for f in args.files)))


if __name__ == "__main__":
    main()
